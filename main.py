"""Extract invoice fields from text-based PDF invoices into an Excel report.

This is a starter project. Validate the extraction rules against each invoice
layout before using it for business-critical accounting.
"""

from __future__ import annotations

import logging
import math
import os
import re
import sys
from collections import Counter
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


# Keep paths stable regardless of the directory from which Python is launched.
BASE_DIR = Path(__file__).resolve().parent
INPUT_FOLDER = BASE_DIR / "input"
OUTPUT_FOLDER = BASE_DIR / "output"
TEXT_FOLDER = OUTPUT_FOLDER / "text"
EXCEL_FILE = OUTPUT_FOLDER / "extracted_invoices.xlsx"
LOG_FILE = OUTPUT_FOLDER / "automation.log"

MAX_LOG_BYTES = 2 * 1024 * 1024  # 2 MiB per log file
LOG_BACKUP_COUNT = 5
MAX_COLUMN_WIDTH = 50

HEADERS = [
    "Invoice Number",
    "Customer",
    "Description",
    "Total Amount",
    "Source PDF",
    "Status",
    "Review Notes",
]


def setup_directories() -> None:
    """Create the input and output directories if they do not exist."""
    for folder in (INPUT_FOLDER, OUTPUT_FOLDER, TEXT_FOLDER):
        folder.mkdir(parents=True, exist_ok=True)


def setup_logging() -> None:
    """Configure console logging and a rotating UTF-8 log file."""
    formatter = logging.Formatter(
        "%(asctime)s [%(levelname)s] %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )

    file_handler = RotatingFileHandler(
        LOG_FILE,
        maxBytes=MAX_LOG_BYTES,
        backupCount=LOG_BACKUP_COUNT,
        encoding="utf-8",
    )
    file_handler.setFormatter(formatter)

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(formatter)

    logging.basicConfig(
        level=logging.INFO,
        handlers=[file_handler, console_handler],
        force=True,
    )


def extract_pdf_text(pdf_path: Path) -> str:
    """Extract and join text from all readable pages in a PDF."""
    reader = PdfReader(str(pdf_path), strict=False)

    if reader.is_encrypted:
        try:
            decrypted = reader.decrypt("")
        except Exception as exc:
            raise ValueError("Unable to decrypt PDF; check whether it needs a password.") from exc
        if not decrypted:
            raise ValueError("PDF is password-protected and no usable password was provided.")

    page_texts: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text()
        except Exception as exc:  # Keep processing other pages when one page is broken.
            logging.warning(
                "Could not extract page %s from %s: %s",
                page_number,
                pdf_path.name,
                exc,
            )
            continue
        if page_text and page_text.strip():
            page_texts.append(page_text.strip())

    return "\n".join(page_texts).strip()


def extract_invoice_number(text: str) -> str | None:
    """Extract invoice identifiers following common labelled patterns."""
    pattern = re.compile(
        r"\bInvoice\s*(?:Number|No\.?|#)\s*[:#\-]?\s*"
        r"([A-Za-z0-9][A-Za-z0-9/_-]*)",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    return match.group(1).strip() if match else None


def extract_customer(text: str) -> str | None:
    """Extract the customer value from a line beginning with Customer:."""
    pattern = re.compile(
        r"^\s*Customer\s*[:\-]\s*(.*?)\s*$",
        re.IGNORECASE | re.MULTILINE,
    )
    match = pattern.search(text)
    if not match:
        return None
    value = match.group(1).strip()
    return value or None


def extract_total_amount(text: str) -> float | None:
    """Extract a labelled total amount; return None when absent or invalid."""
    pattern = re.compile(
        r"\bTotal\s+Amount\s*[:\-]?\s*"
        r"(?:(?:INR|Rs\.?|₹)\s*)?"
        r"([+-]?\d[\d,]*(?:\.\d{1,2})?)",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    if not match:
        return None

    try:
        amount = float(match.group(1).replace(",", ""))
    except ValueError:
        return None
    return amount if math.isfinite(amount) else None


def extract_description(text: str) -> str | None:
    """Extract the line-item description between the table header and totals.

    This rule is designed for the included INV-1001 example and similar layouts.
    Other vendor invoice templates may need different extraction rules.
    """
    header_match = re.search(r"\bDescription\s+Amount\b", text, re.IGNORECASE)
    if not header_match:
        return None

    block = text[header_match.end():]
    summary_match = re.search(
        r"(?im)^\s*(?:Subtotal|Tax(?:\s*\([^\n)]*\))?|"
        r"Grand\s+Total|Total\s+Amount|Balance\s+Due)\b.*$",
        block,
    )
    if summary_match:
        block = block[:summary_match.start()]

    lines = [re.sub(r"\s+", " ", line).strip() for line in block.splitlines()]
    lines = [line for line in lines if line]

    # Remove repeated amount-column headings and a standalone row number.
    while lines and re.fullmatch(r"Amount(?:\s*\([^)]*\))?", lines[0], re.IGNORECASE):
        lines.pop(0)
    if lines and re.fullmatch(r"\d+[.)]?", lines[0]):
        lines.pop(0)

    # Remove a standalone line-item price at the end, not numeric text in a name.
    if lines and re.fullmatch(
        r"(?:(?:INR|Rs\.?|₹)\s*)?[\d,]+(?:\.\d{1,2})?",
        lines[-1],
        re.IGNORECASE,
    ):
        lines.pop()

    description = " ".join(lines).strip()
    return description or None


def save_text_file(pdf_path: Path, text: str) -> Path:
    """Save extracted text for troubleshooting and review."""
    output_path = TEXT_FOLDER / f"{pdf_path.stem}.txt"
    output_path.write_text(text, encoding="utf-8")
    return output_path


def build_record(pdf_file: Path, text: str) -> dict[str, Any]:
    """Extract fields and flag incomplete records for manual review."""
    record: dict[str, Any] = {
        "invoice_number": extract_invoice_number(text),
        "customer": extract_customer(text),
        "description": extract_description(text),
        "total_amount": extract_total_amount(text),
        "source_pdf": pdf_file.name,
        "status": "OK",
        "review_notes": [],
    }

    required_fields = {
        "invoice_number": "Invoice number not found",
        "customer": "Customer not found",
        "description": "Description not found",
        "total_amount": "Total amount not found or invalid",
    }
    for field, message in required_fields.items():
        if record[field] is None:
            record["review_notes"].append(message)

    if record["review_notes"]:
        record["status"] = "REVIEW"
    return record


def flag_duplicate_invoices(records: list[dict[str, Any]]) -> None:
    """Mark duplicate invoice numbers in the current batch for human review."""
    counts = Counter(
        record["invoice_number"].casefold()
        for record in records
        if record["invoice_number"]
    )
    for record in records:
        invoice_number = record["invoice_number"]
        if invoice_number and counts[invoice_number.casefold()] > 1:
            note = "Duplicate invoice number; verify manually"
            if note not in record["review_notes"]:
                record["review_notes"].append(note)
            record["status"] = "REVIEW"


def create_excel_report(records: list[dict[str, Any]], output_path: Path) -> None:
    """Write an Excel report with filters, frozen headers, and review highlighting."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice Report"
    worksheet.append(HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for record in records:
        worksheet.append([
            record["invoice_number"] or "Not Found",
            record["customer"] or "Not Found",
            record["description"] or "Not Found",
            record["total_amount"],
            record["source_pdf"],
            record["status"],
            "; ".join(record["review_notes"]),
        ])
        row_number = worksheet.max_row
        amount_cell = worksheet.cell(row=row_number, column=4)
        if amount_cell.value is not None:
            amount_cell.number_format = '"₹"#,##0.00;[Red]-"₹"#,##0.00'

        if record["status"] == "REVIEW":
            review_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
            for cell in worksheet[row_number]:
                cell.fill = review_fill

    for column in worksheet.columns:
        max_length = max(
            (
                len(line)
                for cell in column
                if cell.value is not None
                for line in str(cell.value).splitlines()
            ),
            default=0,
        )
        column_letter = get_column_letter(column[0].column)
        worksheet.column_dimensions[column_letter].width = min(max_length + 3, MAX_COLUMN_WIDTH)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=cell.column in (3, 7),
            )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30

    # Write a temporary workbook before replacing the last complete report.
    temp_path = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def main() -> int:
    """Process all PDFs in input/ and write one Excel report."""
    setup_directories()
    setup_logging()

    logging.info("=" * 60)
    logging.info("PDF INVOICE AUTOMATION STARTED")
    logging.info("Input folder: %s", INPUT_FOLDER)

    pdf_files = sorted(
        (path for path in INPUT_FOLDER.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"),
        key=lambda path: path.name.casefold(),
    )
    if not pdf_files:
        logging.warning("No PDF files found. Copy a sample PDF into the input folder and rerun.")
        return 0

    records: list[dict[str, Any]] = []
    failed = 0

    for pdf_file in pdf_files:
        logging.info("Processing: %s", pdf_file.name)
        try:
            text = extract_pdf_text(pdf_file)
            if not text:
                logging.warning(
                    "No extractable text in %s. It may be scanned and require OCR.",
                    pdf_file.name,
                )
                failed += 1
                continue

            save_text_file(pdf_file, text)
            record = build_record(pdf_file, text)
            records.append(record)

            if record["status"] == "REVIEW":
                logging.warning(
                    "Review required for %s: %s",
                    pdf_file.name,
                    "; ".join(record["review_notes"]),
                )
            else:
                logging.info(
                    "Extracted %s | %s | INR %.2f",
                    record["invoice_number"],
                    record["customer"],
                    record["total_amount"],
                )
        except Exception:
            failed += 1
            logging.exception("Failed to process %s", pdf_file.name)

    flag_duplicate_invoices(records)
    successful = sum(record["status"] == "OK" for record in records)
    needs_review = sum(record["status"] == "REVIEW" for record in records)

    if records:
        try:
            create_excel_report(records, EXCEL_FILE)
            logging.info("Excel report saved: %s", EXCEL_FILE)
        except PermissionError:
            logging.exception("Cannot replace report. Close it in Excel and try again.")
            return 1
        except Exception:
            logging.exception("Failed to create the Excel report.")
            return 1
    else:
        logging.warning("No invoice text was extracted; no new report was written.")

    logging.info("PDFs found: %d | Passed: %d | Review: %d | Failed/skipped: %d",
                 len(pdf_files), successful, needs_review, failed)
    logging.info("Text backups: %s", TEXT_FOLDER)
    logging.info("Log file: %s", LOG_FILE)
    logging.info("=" * 60)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
